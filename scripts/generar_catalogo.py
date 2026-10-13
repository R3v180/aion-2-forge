"""Genera data/catalogo/<clase>.json (español + inglés) cruzando varias fuentes.

Fuentes:
  - questlog.gg (API tRPC): catálogo completo, nombres oficiales es/en, especializaciones,
    enfriamiento, coste por nivel y valores de daño por nivel.
  - builds de referencia (data/builds/*.js): textos en español ya revisados.
  - Excel del Cantor (opcional, --excel): cotejo de nombres.
  - expcarry.com: cotejo de los nombres en inglés que citan sus guías.

Uso:
  python scripts/generar_catalogo.py [--excel RUTA.xlsx] [--refrescar]
"""
import argparse
import html
import json
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
FUENTES = RAIZ / "scripts" / "fuentes"
SALIDA = RAIZ / "data" / "catalogo"
ICONOS = RAIZ / "icons" / "habilidades"
API = "https://questlog.gg/aion-2/api/trpc/"
CDN = "https://cdn.questlog.gg/aion-2"
UA = {"User-Agent": "Mozilla/5.0 (BuildsAion2 catalog generator)"}

# id interno -> categoría de questlog
CLASES = {
    "gladiator": "gladiator",
    "templar": "templar",
    "assassin": "assassin",
    "ranger": "ranger",
    "sorcerer": "sorcerer",
    "spiritmaster": "elementalist",
    "cleric": "cleric",
    "chanter": "chanter",
}
EXPCARRY = {c: f"https://expcarry.com/aion-2-{c}-guide" for c in CLASES}
TIPOS = {"active": "activa", "stigma": "estigma", "passive": "pasiva"}


def pedir(url, intentos=3):
    for i in range(intentos):
        try:
            return urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=90).read()
        except (urllib.error.URLError, TimeoutError) as e:
            if i == intentos - 1:
                raise
            time.sleep(2 + i * 3)


def trpc(proc, entrada, cache, refrescar):
    if cache.exists() and not refrescar:
        return json.loads(cache.read_text(encoding="utf-8"))
    url = API + proc + "?input=" + urllib.parse.quote(json.dumps(entrada))
    datos = json.loads(pedir(url))["result"]["data"]
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(json.dumps(datos, ensure_ascii=False), encoding="utf-8")
    time.sleep(0.25)
    return datos


def limpiar(texto):
    """HTML de questlog -> texto plano (los marcadores {…} se conservan)."""
    t = re.sub(r"<br\s*/?>", "\n", texto or "")
    t = re.sub(r"<[^>]+>", "", t)
    return html.unescape(t).replace("\u00a0", " ").strip()


def numero(v):
    return float(v) if v not in (None, "") and re.fullmatch(r"-?\d+(\.\d+)?", str(v)) else None


def tabla_dano(ph):
    """Marcador se_dmg -> [[%ATQ, fijo, golpes], …] por nivel (índice 0 = nivel 1)."""
    filas = []
    base = ph.get("base") or {}
    niveles = ph.get("levels") or {}
    tope = max([1] + [int(k) for k in niveles])
    for n in range(1, tope + 1):
        v = (niveles.get(str(n)) or base).get("values") or []
        fijo, pct, golpes = numero(v[0] if v else None), numero(v[2] if len(v) > 2 else None), numero(v[4] if len(v) > 4 else None)
        if fijo is None or pct is None:
            return None
        filas.append([round(pct / 100, 2), int(fijo), int(golpes or 1)])
    return filas


def describir(desc_render, desc_bruta, placeholders):
    """Sustituye cada par {min}-{max} de daño por {dN}; devuelve texto y tablas de daño."""
    texto = limpiar(desc_render or desc_bruta)
    tablas = {}
    for clave, ph in (placeholders or {}).items():
        if ph.get("type") != "se_dmg" or "min" not in clave.lower():
            continue
        t = tabla_dano(ph)
        if not t:
            continue
        par = re.escape(clave) + r"\s*[-~]\s*" + re.escape(clave.replace("Min", "Max").replace("min", "max"))
        nombre = "d%d" % len(tablas)
        nuevo, n = re.subn(par, "{" + nombre + "}", texto)
        if not n:
            nuevo, n = re.subn(re.escape(clave), "{" + nombre + "}", texto)
        if n:
            texto = nuevo
            tablas[nombre] = t
    # daño periódico: questlog no publica su base, el juego tampoco lo muestra fuera de combate
    marcador = r"\{(?:se|abe|sef|se_abe|se_dmg|se_abe_dmg):[^}]+\}"
    texto = re.sub(marcador + r"\s*[-~]\s*" + marcador, "—", texto)
    texto = re.sub(marcador, "—", texto)
    return texto, tablas


def icono_local(ruta):
    if not ruta:
        return None
    nombre = ruta.rsplit("/", 1)[-1].split(".")[0]
    destino = ICONOS / (nombre + ".webp")
    if not destino.exists():
        try:
            destino.write_bytes(pedir(f"{CDN}{ruta.rsplit('.', 1)[0]}.webp"))
            print("  icono descargado:", destino.name)
        except Exception as e:  # noqa: BLE001
            print("  icono NO disponible:", nombre, e)
            return None
    return nombre


def leer_referencia(clase):
    f = RAIZ / "data" / "builds" / f"{clase}.js"
    if not f.exists():
        return {}
    s = f.read_text(encoding="utf-8")
    b = json.loads(s[s.index("{"): s.rindex("}") + 1])
    return {h["id"]: h for h in b["habilidades"]}


def leer_excel(ruta):
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True)
    filas = list(wb.worksheets[0].iter_rows(values_only=True))[1:]
    return {str(r[0]).strip() for r in filas if r and r[0]}


def nombres_expcarry(clase, refrescar):
    cache = FUENTES / "expcarry" / f"{clase}.html"
    if not cache.exists() or refrescar:
        try:
            cache.parent.mkdir(parents=True, exist_ok=True)
            cache.write_bytes(pedir(EXPCARRY[clase]))
        except Exception as e:  # noqa: BLE001
            return None, str(e)
    return re.sub(r"<[^>]+>", " ", cache.read_text(encoding="utf-8", errors="replace")), None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", help="Excel con habilidades (cotejo de nombres en español)")
    ap.add_argument("--refrescar", action="store_true", help="vuelve a descargar todo de questlog")
    args = ap.parse_args()

    SALIDA.mkdir(parents=True, exist_ok=True)
    ICONOS.mkdir(parents=True, exist_ok=True)
    lista = {}
    for lang in ("es", "en"):
        lista[lang] = {s["id"]: s for s in trpc("skillBuilder.getSkills", {"language": lang},
                                                 FUENTES / "questlog" / f"skills_{lang}.json", args.refrescar)}
    excel = leer_excel(args.excel) if args.excel else set()
    informe = []

    for clase, cat in CLASES.items():
        print(f"== {clase}")
        zen = leer_referencia(clase)
        ids = sorted(i for i, s in lista["es"].items() if s.get("mainCategory") == cat and s["type"] in TIPOS)
        habilidades = []
        for sid in ids:
            fichas = {lang: trpc("database.getSkill", {"id": sid, "language": lang},
                                 FUENTES / "questlog" / "skill" / f"{sid}_{lang}.json", args.refrescar) for lang in ("es", "en")}
            base = lista["es"][sid]
            h = {"id": sid, "tipo": TIPOS[base["type"]], "icono": icono_local(base.get("icon")),
                 "nombre": {}, "desc": {}, "dano": {}, "esp": []}
            for lang in ("es", "en"):
                s, f = lista[lang][sid], fichas[lang]
                h["nombre"][lang] = s["name"]
                texto, tablas = describir(f.get("description"), s.get("description"),
                                          (s.get("descriptionData") or {}).get("placeholders"))
                h["desc"][lang] = texto
                if lang == "es":
                    h["dano"] = tablas
            f = fichas["es"]
            h["enfriamiento"] = round((f.get("cooldown") or 0) / 1000, 1) or None
            h["armas"] = f.get("requiredWeapons") or []
            h["coste"] = [[c["level"], c.get("skillPoints") or 0, c.get("stigmaPoints") or 0, c.get("characterLevel")]
                          for c in (f.get("learning") or [])]
            esp_en = lista["en"][sid].get("specializations") or []
            for n, sp in enumerate(lista["es"][sid].get("specializations") or [], 1):
                en = esp_en[n - 1] if n - 1 < len(esp_en) else {}
                h["esp"].append({"n": n, "nivel": sp.get("parentSkillLvl"),
                                 "texto": {"es": limpiar(sp.get("specialized")), "en": limpiar(en.get("specialized"))}})
            # cotejo con la referencia
            z = zen.get(sid)
            if z:
                if z["nombre"] != h["nombre"]["es"]:
                    # questlog repite algún nombre en español; el de la referencia está revisado en el juego
                    informe.append(f"{clase}: nombre distinto en la referencia «{z['nombre']}» vs questlog «{h['nombre']['es']}» (se usa el de la referencia)")
                    h["nombre"]["es"] = z["nombre"]
                h["fuentes"] = ["questlog", "referencia"]
            else:
                h["fuentes"] = ["questlog"]
            habilidades.append(h)

        # cotejo: la referencia tiene algo que questlog no
        for zid, z in zen.items():
            if zid not in ids:
                informe.append(f"{clase}: «{z['nombre']}» ({zid}) está en la referencia pero no en questlog")
        # cotejo: Excel (solo Cantor)
        if clase == "chanter" and excel:
            es_q = {h["nombre"]["es"] for h in habilidades}
            for n in sorted(excel - es_q):
                informe.append(f"{clase}: «{n}» está en el Excel pero no en questlog")
            for n in sorted(es_q - excel):
                informe.append(f"{clase}: «{n}» está en questlog pero no en el Excel")
        # cotejo: expcarry cita habilidades en inglés
        texto, err = nombres_expcarry(clase, args.refrescar)
        if texto is None:
            informe.append(f"{clase}: expcarry no disponible ({err})")
        else:
            citadas = [h["nombre"]["en"] for h in habilidades if h["nombre"]["en"] and h["nombre"]["en"] in texto]
            informe.append(f"{clase}: expcarry cita {len(citadas)}/{len(habilidades)} habilidades del catálogo")

        cuenta = {t: sum(1 for h in habilidades if h["tipo"] == t) for t in TIPOS.values()}
        print("  ", cuenta, "| con referencia:", sum(1 for h in habilidades if "referencia" in h["fuentes"]))
        (SALIDA / f"{clase}.json").write_text(json.dumps({"clase": clase, "habilidades": habilidades},
                                                         ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    (FUENTES / "informe.txt").write_text("\n".join(informe) + "\n", encoding="utf-8")
    print("\nInforme de cotejo:")
    print("\n".join(informe))


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
